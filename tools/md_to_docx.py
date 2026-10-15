"""把 docs 下的设计文档（Markdown）转换为 Word（.docx）。

用法::

    python tools/md_to_docx.py docs/每日锻炼计划功能规划.md docs/每日锻炼计划功能规划.docx

支持的 Markdown 子集：标题、段落、无序/有序列表、引用、围栏代码块、表格、行内
加粗与行内代码。仅依赖 python-docx。
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

LATIN_FONT = "Calibri"
CJK_FONT = "微软雅黑"
MONO_FONT = "Consolas"
HEADER_FILL = "D9E2F3"


def set_cjk(run, font_name: str) -> None:
    """同时设置西文字体与中文字体（东亚字体需单独指定 xml 属性）。"""
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def configure_document(doc: Document) -> None:
    """设置全局默认字体与页面尺寸。"""
    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    style.font.name = LATIN_FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    for section in doc.sections:
        section.page_width, section.page_height = Pt(595.3), Pt(841.9)  # A4
        for attr in ("top_margin", "bottom_margin"):
            setattr(section, attr, Pt(56.7))
        for attr in ("left_margin", "right_margin"):
            setattr(section, attr, Pt(70.9))


def shade_cell(cell, color: str = HEADER_FILL) -> None:
    """给单元格填充底色。"""
    from docx.oxml import OxmlElement

    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(el)


def add_inline(paragraph, text: str, base_bold: bool = False) -> None:
    """解析行内 ``**加粗**`` 与 ``` `代码` ``` 并写入段落。"""
    for chunk in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(9.5)
        else:
            run = paragraph.add_run(chunk)
            run.bold = base_bold
        set_cjk(run, CJK_FONT if not run.font.name == MONO_FONT else MONO_FONT)


def add_code_block(doc: Document, lines: list[str]) -> None:
    """插入等宽字体的代码块（单单元格表格以保留底色）。"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "F2F2F2")
    cell.text = ""
    for i, line in enumerate(lines):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        run = para.add_run(line)
        run.font.name = MONO_FONT
        run.font.size = Pt(8.5)
        para.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    """插入表格，首行为表头。"""
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if r == 0 else WD_ALIGN_PARAGRAPH.LEFT
            stripped = text.replace("\\|", "|")
            add_inline(para, stripped, base_bold=(r == 0))
            if r == 0:
                shade_cell(cell)
            for run in para.runs:
                run.font.size = Pt(9)
    doc.add_paragraph()


def split_table_row(line: str) -> list[str]:
    """拆分 Markdown 表格行。"""
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    """判断是否为表格分隔行（|---|---|）。"""
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def convert(md_path: Path, docx_path: Path) -> None:
    """执行转换。"""
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):                              # 代码块
            i += 1
            block: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, block)
        elif stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            rows = [split_table_row(stripped)]                       # 表格
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        elif stripped.startswith("#"):                               # 标题
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1:
                para = doc.add_heading(level=0)
                add_inline(para, text)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(20)
                    run.font.color.rgb = RGBColor(0, 0, 0)
            else:
                para = doc.add_heading(level=min(level - 1, 4))
                add_inline(para, text)
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
        elif stripped.startswith("- "):                              # 无序列表
            para = doc.add_paragraph(style="List Bullet")
            add_inline(para, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):                        # 有序列表
            para = doc.add_paragraph(style="List Number")
            add_inline(para, re.sub(r"^\d+\.\s", "", stripped))
        elif stripped.startswith(">"):                               # 引用
            para = doc.add_paragraph(style="Intense Quote")
            add_inline(para, stripped.lstrip("> ").strip())
        elif stripped in ("---", "***"):                             # 分隔线
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(6)
            run = para.add_run("─" * 40)
            run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
        elif stripped:                                               # 普通段落
            para = doc.add_paragraph()
            add_inline(para, stripped)
        i += 1

    doc.save(str(docx_path))
    print(f"已生成: {docx_path}")


def main() -> None:
    """命令行入口。"""
    if len(sys.argv) != 3:
        print("用法: python tools/md_to_docx.py <input.md> <output.docx>")
        raise SystemExit(1)
    convert(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
